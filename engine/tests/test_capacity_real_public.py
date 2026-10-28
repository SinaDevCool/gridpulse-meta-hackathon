from __future__ import annotations

import json
import os
import zipfile
from pathlib import Path

from openpyxl import Workbook

from capacity_backtest.capacity_study.annual import execute_resumable, plan_annual_chunks
from capacity_backtest.capacity_study.enquiry import build_enquiry_package
from capacity_backtest.capacity_study.fifty_hertz_static import (
    normalize_workbook,
    profile_workbook,
    write_normalized,
)
from capacity_backtest.capacity_study.mastr_assets import (
    DownloadLockError,
    _acquire_lock,
    _validate_checkpoint,
    _write_checkpoint,
    _zip_inventory,
    normalize_mastr_archive,
)
from capacity_backtest.capacity_study.validation import ReconstructionCase, evaluate_reconstruction


def workbook(path: Path) -> None:
    book = Workbook()
    lines = book.active
    lines.title = "Lines"
    lines.append([None] * 24)
    lines.append([None] * 24)
    row = [None] * 24
    row[:8] = ["A-B", "line-1", "50HERTZ", "A", "A1", "B", "B1", 380]
    row[14], row[17], row[18], row[19], row[20] = 1000, 1.0, 5.0, 100.0, 10.0
    lines.append(row)
    ties = book.create_sheet("Tielines")
    ties.append([None] * 24)
    ties.append([None] * 24)
    tie = [None] * 24
    tie[:9] = ["A-X", "tie-1", "50HERTZ", "A", "A1", "OTHER", "X", "X1", 380]
    tie[15], tie[18], tie[19], tie[20], tie[21] = 900, 1.0, 4.0, 90.0, 8.0
    ties.append(tie)
    transformers = book.create_sheet("Transformers")
    transformers.append([None] * 27)
    transformers.append([None] * 27)
    remedial = book.create_sheet("Remedial Actions")
    remedial.append([None] * 3)
    remedial.append([None] * 3)
    remedial.append(["PST A", "50HERTZ", "Change tap position"])
    book.save(path)


def test_50hertz_profile_and_normalization(tmp_path: Path) -> None:
    path = tmp_path / "network.xlsx"
    workbook(path)
    assert {item["name"] for item in profile_workbook(path)["sheets"]} == {
        "Lines",
        "Tielines",
        "Transformers",
        "Remedial Actions",
    }
    model, quality = normalize_workbook(path)
    assert quality["line_count"] == 2
    assert quality["boundary_candidates"][0]["remote_operator"] == "OTHER"
    assert quality["remedial_actions"][0]["availability"].endswith("dispatch_permission")
    assert model.branches[0]["r_ohm_per_km"] == 0.1
    destination = write_normalized(model, quality, tmp_path / "data")
    assert (destination / "boundary-candidates.json").is_file()
    assert (destination / "remedial-actions.json").is_file()


def test_mastr_streaming_region_filter(tmp_path: Path) -> None:
    archive = tmp_path / "mastr.zip"
    xml = b"<Root><Einheit><EinheitMastrNummer>SEE1</EinheitMastrNummer><Postleitzahl>10115</Postleitzahl><Nettonennleistung>10</Nettonennleistung></Einheit></Root>"
    with zipfile.ZipFile(archive, "w") as handle:
        handle.writestr("assets.xml", xml)
    destination = normalize_mastr_archive(archive, tmp_path / "data")
    assert "SEE1" in (destination / "assets.jsonl").read_text()


def test_mastr_download_lock_rejects_live_owner_and_reclaims_stale(tmp_path: Path) -> None:
    lock = tmp_path / "mastr.lock"
    _acquire_lock(lock, {"pid": os.getpid()})
    try:
        try:
            _acquire_lock(lock, {"pid": os.getpid()})
        except DownloadLockError:
            pass
        else:
            raise AssertionError("live lock was not rejected")
    finally:
        lock.unlink()
    lock.write_text(json.dumps({"pid": 99999999}))
    _acquire_lock(lock, {"pid": os.getpid()})
    assert json.loads(lock.read_text())["pid"] == os.getpid()


def test_mastr_checkpoint_detects_remote_identity_change(tmp_path: Path) -> None:
    staging = tmp_path / "mastr.zip"
    staging.write_bytes(b"partial")
    values = {
        "source_url": "https://example.test/mastr.zip",
        "snapshot_date": "2025-01-01",
        "expected_size": 100,
        "etag": "fixed",
        "last_modified": "date",
    }
    _write_checkpoint(staging, **values)
    _validate_checkpoint(staging, **values)
    values["etag"] = "changed"
    try:
        _validate_checkpoint(staging, **values)
    except ValueError:
        pass
    else:
        raise AssertionError("changed remote identity was not rejected")


def test_mastr_zip_inventory_checks_expected_families(tmp_path: Path) -> None:
    archive = tmp_path / "full.zip"
    with zipfile.ZipFile(archive, "w") as handle:
        handle.writestr("EinheitenSolar.xml", "<x/>")
        handle.writestr("Marktakteure.xml", "<x/>")
        handle.writestr("Lokationen.xml", "<x/>")
    inventory = _zip_inventory(archive)
    assert len(inventory) == 3


def test_resumable_ledger_and_validation(tmp_path: Path) -> None:
    worker = lambda item: {"value": item["value"] * 2}
    assert execute_resumable([{"value": 2}], tmp_path, worker)["completed"] == 1
    assert execute_resumable([{"value": 2}], tmp_path, worker)["reused"] == 1
    metrics = evaluate_reconstruction([ReconstructionCase("a", 10, 8, 10, 12, "line", "line")])
    assert metrics["release_gate_passed"]


def test_compact_annual_chunk_plan(tmp_path: Path) -> None:
    rows = tmp_path / "rows.jsonl"
    rows.write_text(
        "".join(
            json.dumps({"timestamp_local": f"2025-01-{day:02d}T00:00:00"}) + "\n"
            for day in range(1, 11)
        )
    )
    plan = json.loads(plan_annual_chunks(rows, tmp_path / "run", hours_per_chunk=4).read_text())
    assert plan["hour_count"] == 10
    assert plan["chunk_count"] == 3


def test_enquiry_is_generated_from_nonclaim_run(tmp_path: Path) -> None:
    run = tmp_path / "run.json"
    run.write_text(
        json.dumps(
            {
                "run_id": "r1",
                "capacity_claim": False,
                "truth_class": "experimental_estimate",
                "configuration": {"site": {"site_id": "s", "requested_capacity_mw": 40}},
                "pccs": [],
                "statistics": {},
            }
        )
    )
    destination = build_enquiry_package(run, tmp_path / "enquiry")
    assert (destination / "enquiry.json").is_file()
