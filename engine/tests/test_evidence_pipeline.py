from __future__ import annotations

import json
from dataclasses import replace
from datetime import datetime
from pathlib import Path

import httpx
import openpyxl
import pytest

from capacity_backtest.allocation import allocation_ratio, parse_berlin_2025_allocation
from capacity_backtest.audit import build_audit, write_audit
from capacity_backtest.contracts import TruthClass, assert_capacity_claim_allowed
from capacity_backtest.download import fetch_source, verify_snapshot
from capacity_backtest.identity import Node, identity_review_queue, propose_match
from capacity_backtest.load_data import normalize_power, parse_load_workbook
from capacity_backtest.registry import load_registry
from capacity_backtest.sufficiency import berlin_public_evidence_assessment

ROOT = Path(__file__).parents[1]


def test_similarly_named_assets_and_collisions_require_review() -> None:
    official_a = Node("a", "Umspannwerk Nord", "E.DIS Netz", 110, 52.5, 13.4, "official")
    official_b = Node("b", "Umspannwerk Nord 2", "E.DIS Netz", 110, 52.51, 13.4, "official")
    candidate = Node("candidate", "UW Nord", "E.DIS", 110, 52.5, 13.4, "public")
    proposals = [propose_match(official_a, candidate), propose_match(official_b, candidate)]
    review = identity_review_queue(proposals)
    assert len(review) == 2
    assert all(item["collision"] for item in review)


def test_registry_is_valid_and_unique() -> None:
    sources = load_registry(ROOT / "config" / "sources")
    assert len(sources) >= 16
    assert len({source.id for source in sources}) == len(sources)
    assert all(source.prohibited_uses for source in sources)
    ids = {source.id for source in sources}
    assert {
        "dwd_open_data",
        "entsoe_transparency",
        "tennet_transparency",
        "amprion_transparency",
        "wesernetz_publications",
        "edis_publications",
    } <= ids


def test_public_and_operator_published_sources_cannot_claim_capacity() -> None:
    for truth_class in (
        TruthClass.PUBLIC_SCREENING,
        TruthClass.OPERATOR_PUBLISHED,
        TruthClass.CALCULATED_VALIDATED,
        TruthClass.OPERATOR_RECONCILED,
        TruthClass.OPERATOR_MODEL_UNVALIDATED,
        TruthClass.SYNTHETIC_BENCHMARK,
    ):
        with pytest.raises(ValueError, match="cannot establish"):
            assert_capacity_claim_allowed(truth_class, 20)
    assert_capacity_claim_allowed(TruthClass.OPERATOR_CONFIRMED, 20)
    assert_capacity_claim_allowed(TruthClass.PUBLIC_SCREENING, None)


def test_downloader_writes_hash_verified_immutable_snapshot(tmp_path: Path) -> None:
    source = load_registry(ROOT / "config" / "sources")[0]
    request = httpx.Request("GET", source.url)
    response = httpx.Response(
        200, content=b"authoritative page", request=request, headers={"content-type": "text/html"}
    )

    class Client:
        def get(self, *_args, **_kwargs):
            return response

    result = fetch_source(source, tmp_path, client=Client())
    assert result.status == "succeeded"
    assert result.manifest_path and verify_snapshot(result.manifest_path)
    manifest = json.loads(result.manifest_path.read_text(encoding="utf-8"))
    assert manifest["sha256"]
    assert manifest["truth_class"] == source.truth_class.value


def test_binary_source_rejects_html_error_page(tmp_path: Path) -> None:
    source = replace(load_registry(ROOT / "config" / "sources")[0], format="pdf")
    request = httpx.Request("GET", source.url)
    response = httpx.Response(
        200,
        content=b"<!doctype html>blocked",
        request=request,
        headers={"content-type": "text/html"},
    )

    class Client:
        def get(self, *_args, **_kwargs):
            return response

    assert fetch_source(source, tmp_path, client=Client()).status == "failed"


def test_load_workbook_preserves_mw_and_flags_gaps(tmp_path: Path) -> None:
    path = tmp_path / "load.xlsx"
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.append(["timestamp", "MW"])
    sheet.append([datetime.fromisoformat("2025-01-01T00:00:00"), 100])
    sheet.append([datetime.fromisoformat("2025-01-01T00:15:00"), 110])
    sheet.append([datetime.fromisoformat("2025-01-01T01:00:00"), 120])
    workbook.save(path)
    validation = parse_load_workbook(path, voltage_level="110 kV")
    assert validation.valid
    assert [point.value for point in validation.points] == [100, 110, 120]
    assert validation.warnings == ["time series contains missing intervals"]


def test_units_never_silently_mix_real_apparent_power_or_energy() -> None:
    assert normalize_power(1000, "kW") == 1
    with pytest.raises(ValueError, match="apparent-power"):
        normalize_power(20, "MVA")
    with pytest.raises(ValueError, match="energy unit"):
        normalize_power(20, "MWh")


def test_allocation_history_does_not_infer_remaining_capacity() -> None:
    text = (
        "Rund 2300 MVA wurden angefragt. 70 Anfragen lagen vor. "
        "47 Anfragen vollständig. 37 Antragstellern konnten 365 MVA vergeben werden."
    )
    round_ = parse_berlin_2025_allocation(text)
    assert round_.requested_mva == 2300
    assert round_.request_count == 70
    assert round_.eligible_request_count == 47
    assert round_.allocated_request_count == 37
    assert round_.allocated_mva == 365
    assert round_.remaining_mva is None
    assert allocation_ratio(round_) == pytest.approx(365 / 2300)


def test_crosswalk_proposals_require_review() -> None:
    official = Node("official", "UW Mitte", "Stromnetz Berlin", 110, 52.52, 13.40, "operator")
    candidate = Node("mapped", "Umspannwerk Mitte", "Stromnetz Berlin", 110, 52.521, 13.401, "map")
    match = propose_match(official, candidate)
    assert match.status != "accepted"
    assert "explicit review" in match.reviewer_note


def test_sufficiency_fails_closed_for_capacity_bands() -> None:
    assessments = {item.output: item for item in berlin_public_evidence_assessment()}
    capacity = assessments["capacity_bands"]
    assert capacity.status.value == "insufficient"
    assert "operator_validation_or_outcomes" in capacity.missing_inputs
    assert "contracted_queued_demand" in capacity.missing_inputs


def test_audit_is_machine_and_human_readable(tmp_path: Path) -> None:
    audit = build_audit(ROOT / "config" / "sources")
    assert "insufficient" in str(audit["decision"]).lower()
    json_path = tmp_path / "audit.json"
    markdown_path = tmp_path / "audit.md"
    write_audit(ROOT / "config" / "sources", json_path, markdown_path)
    assert json.loads(json_path.read_text(encoding="utf-8"))["study_area"] == "Berlin"
    assert "Unknown information remains unknown" in markdown_path.read_text(encoding="utf-8")
