from datetime import datetime, time
from io import BytesIO

import openpyxl

from capacity_backtest.regional_load import (
    profile_berlin_load_workbook,
    profile_wesernetz_load_csv,
)


def test_berlin_workbook_preserves_kw_and_rejects_unspecified_timezone() -> None:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "2025"
    sheet.append(["Max in kW", None, 2000])
    sheet.append(["Datum", "Zeit", "Wert"])
    sheet.append([datetime(2025, 1, 1), time(0, 15), 1000])  # noqa: DTZ001
    buffer = BytesIO()
    workbook.save(buffer)
    report = profile_berlin_load_workbook(
        buffer.getvalue(), source_id="berlin", snapshot_hash="a" * 64
    )
    assert report["total_points"] == 1
    assert report["units_validated"] is True
    assert report["timezone_validated"] is False
    assert report["feature_admitted"] is False


def test_wesernetz_profiles_dst_duplicates_without_guessing_fold() -> None:
    content = "Zeit;MW\n26.10.2025 02:00;100.5\n26.10.2025 02:00;101.5\n".encode("cp1252")
    report = profile_wesernetz_load_csv(content, source_id="wesernetz", snapshot_hash="b" * 64)
    assert report["record_count"] == 2
    assert report["duplicate_local_timestamps"] == 1
    assert report["dst_semantics_validated"] is False
    assert report["feature_admitted"] is False
