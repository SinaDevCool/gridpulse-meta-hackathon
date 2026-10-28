"""Capture, profile and normalize the official 50Hertz static-model workbook."""

from __future__ import annotations

import hashlib
import json
import math
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook  # type: ignore[import-untyped]

from capacity_backtest.storage import SnapshotStore, atomic_write
from grid_data.network_study import NetworkModelInput

from .network_adapter import validate_network_model

SOURCE_ID = "50hertz-static-network-model"
SOURCE_URL = "https://www.50hertz.com/cdn/files/c62467c9-a504-4be0-c6ee-08debc9147be/statisches%20netzmodell%20-%20datentabelle%202025-q4%20-%20oo.xlsx"
SOURCE_PAGE = "https://www.50hertz.com/de/Transparenz/Kennzahlen/Netzdaten/StatischesNetzmodell"
LICENCE = "50Hertz published transparency data; reuse terms require publisher review"
OMISSIONS = (
    "Customer and power-plant connection lines are generally omitted.",
    "Non-meshed spur circuits are omitted.",
    "Some 380/220-kV transformers feeding islands or third parties are omitted.",
    "The workbook is a transmission model, not the Stromnetz Berlin DSO model.",
)


def capture_static_model(data_root: Path, work_root: Path, *, source_url: str = SOURCE_URL) -> Path:
    staging = work_root / "staging" / "50hertz-static-network-model.xlsx"
    staging.parent.mkdir(parents=True, exist_ok=True)
    with httpx.stream("GET", source_url, follow_redirects=True, timeout=120) as response:
        response.raise_for_status()
        content = response.read()
        headers = dict(response.headers)
    if not content.startswith(b"PK"):
        raise ValueError("50Hertz static-model response is not an XLSX container")
    atomic_write(staging, content)
    stored = SnapshotStore(data_root).capture(
        source_id=SOURCE_ID,
        publisher="50Hertz Transmission GmbH",
        source_url=source_url,
        content=staging.read_bytes(),
        extension="xlsx",
        retrieved_at=datetime.now(timezone.utc),
        response_headers=headers,
        truth_class="operator_published",
        geographic_resolution="380/220-kV meshed transmission assets",
        licence=LICENCE,
        limitations=OMISSIONS,
        adapter_version="50hertz-static-v1",
    )
    staging.unlink(missing_ok=True)
    return stored.manifest_path


def profile_workbook(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=False, data_only=False)
    sheets = []
    for sheet in workbook.worksheets:
        non_empty_rows = []
        formula_count = 0
        for row_index, row in enumerate(sheet.iter_rows(), start=1):
            values = [cell.value for cell in row]
            formula_count += sum(
                isinstance(value, str) and value.startswith("=") for value in values
            )
            if any(value is not None for value in values):
                non_empty_rows.append(row_index)
        headers = [cell.value for cell in sheet[2]] if sheet.max_row >= 2 else []
        sheets.append(
            {
                "name": sheet.title,
                "max_row": sheet.max_row,
                "max_column": sheet.max_column,
                "non_empty_row_count": len(non_empty_rows),
                "first_non_empty_row": min(non_empty_rows, default=None),
                "last_non_empty_row": max(non_empty_rows, default=None),
                "formula_count": formula_count,
                "merged_range_count": len(sheet.merged_cells.ranges),
                "hidden": sheet.sheet_state != "visible",
                "headers": [str(value) if value is not None else None for value in headers],
            }
        )
    result = {
        "schema_version": "gridpulse-workbook-profile-v1",
        "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
        "sheets": sheets,
    }
    workbook.close()
    return result


def _rows(sheet) -> list[tuple[Any, ...]]:
    return [
        row
        for row in sheet.iter_rows(min_row=3, values_only=True)
        if any(value is not None for value in row)
    ]


def normalize_workbook(path: Path) -> tuple[NetworkModelInput, dict[str, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    buses: dict[str, dict[str, Any]] = {}
    branches = []
    boundary_candidates = []
    asset_ids: set[str] = set()
    assumptions = []
    for sheet_name in ("Lines", "Tielines"):
        for row in _rows(workbook[sheet_name]):
            if len(row) < 22 or row[7 if sheet_name == "Lines" else 8] is None:
                continue
            if sheet_name == "Lines":
                name, eic, tso, full_a, short_a, full_b, short_b, voltage = row[:8]
                fixed, resistance, reactance, susceptance, length = (
                    row[14],
                    row[17],
                    row[18],
                    row[19],
                    row[20],
                )
            else:
                name, eic, tso, full_a, short_a, tso_b, full_b, short_b, voltage = row[:9]
                fixed, resistance, reactance, susceptance, length = (
                    row[15],
                    row[18],
                    row[19],
                    row[20],
                    row[21],
                )
                boundary_candidates.append(
                    {
                        "asset_id": str(eic or name),
                        "local_bus": str(short_a),
                        "remote_bus": str(short_b),
                        "local_operator": str(tso),
                        "remote_operator": str(tso_b),
                        "remote_station": str(full_b),
                        "voltage_kv": float(voltage),
                        "status": "published_tie_line_candidate_not_operating_boundary",
                    }
                )
            if (
                not short_a
                or not short_b
                or not length
                or not fixed
                or resistance is None
                or reactance is None
            ):
                continue
            for short, full in ((short_a, full_a), (short_b, full_b)):
                buses.setdefault(
                    str(short), {"id": str(short), "name": str(full), "vn_kv": float(voltage)}
                )
            length = float(length)
            c_nf_per_km = float(susceptance or 0) * 1000 / (2 * math.pi * 50 * length)
            source_asset_id = str(eic or name)
            asset_id = source_asset_id
            suffix = 2
            while asset_id in asset_ids:
                asset_id = f"{source_asset_id}-{sheet_name.lower()}-{suffix}"
                suffix += 1
            asset_ids.add(asset_id)
            branches.append(
                {
                    "id": asset_id,
                    "source_eic": source_asset_id,
                    "name": str(name),
                    "from_bus": str(short_a),
                    "to_bus": str(short_b),
                    "length_km": length,
                    "r_ohm_per_km": float(resistance) / length,
                    "x_ohm_per_km": float(reactance) / length,
                    "c_nf_per_km": c_nf_per_km,
                    "max_i_ka": float(fixed) / 1000,
                    "source_sheet": sheet_name,
                    "operator": str(tso),
                }
            )
    transformers = []
    for row in _rows(workbook["Transformers"]):
        if len(row) < 13 or row[7] is None or row[8] is None:
            continue
        (
            name,
            short,
            eic,
            tso,
            minimum_a,
            maximum_a,
            fixed_a,
            primary,
            secondary,
            resistance,
            reactance,
        ) = row[:11]
        primary, secondary = float(primary), float(secondary)
        hv_bus = str(short)
        candidate_secondary = f"{str(short)[:-1]}2" if str(short).endswith("1") else ""
        secondary_bus = (
            candidate_secondary if candidate_secondary in buses else f"{short}-lv-{secondary:g}"
        )
        buses.setdefault(hv_bus, {"id": hv_bus, "name": str(name), "vn_kv": primary})
        buses.setdefault(
            secondary_bus, {"id": secondary_bus, "name": f"{name} secondary", "vn_kv": secondary}
        )
        current_a = float(fixed_a or maximum_a or minimum_a or 0)
        if current_a <= 0:
            continue
        sn_mva = math.sqrt(3) * primary * current_a / 1000
        z_base = primary * primary / sn_mva
        r_value, x_value = float(resistance or 0), float(reactance or 0)
        source_asset_id = str(eic or name)
        asset_id = source_asset_id
        suffix = 2
        while asset_id in asset_ids:
            asset_id = f"{source_asset_id}-transformer-{suffix}"
            suffix += 1
        asset_ids.add(asset_id)
        transformers.append(
            {
                "id": asset_id,
                "source_eic": source_asset_id,
                "name": str(name),
                "hv_bus": hv_bus,
                "lv_bus": secondary_bus,
                "sn_mva": sn_mva,
                "vn_hv_kv": primary,
                "vn_lv_kv": secondary,
                "vk_percent": 100 * math.hypot(r_value, x_value) / z_base,
                "vkr_percent": 100 * r_value / z_base,
                "pfe_kw": 0.0,
                "i0_percent": 0.0,
                "operator": str(tso),
                "secondary_bus_inferred": True,
            }
        )
        if secondary_bus.endswith(f"lv-{secondary:g}"):
            assumptions.append(f"Transformer secondary bus inferred for {eic or name}")
    if not buses or not branches:
        raise ValueError("50Hertz workbook produced no usable buses or lines")
    remedial_actions = [
        {
            "name": str(row[0]),
            "operator": str(row[1]),
            "description": str(row[2]),
            "availability": "published_capability_not_dispatch_permission",
        }
        for row in _rows(workbook["Remedial Actions"])
        if len(row) >= 3 and row[0]
    ]
    adjacency: dict[str, set[str]] = {bus_id: set() for bus_id in buses}
    for item in branches:
        from_bus, to_bus = str(item["from_bus"]), str(item["to_bus"])
        adjacency[from_bus].add(to_bus)
        adjacency[to_bus].add(from_bus)
    for item in transformers:
        hv_bus, lv_bus = str(item["hv_bus"]), str(item["lv_bus"])
        adjacency[hv_bus].add(lv_bus)
        adjacency[lv_bus].add(hv_bus)
    component_slacks = []
    unseen = set(buses)
    while unseen:
        slack = min(unseen)
        stack, component = [slack], set()
        while stack:
            current = stack.pop()
            if current in component:
                continue
            component.add(current)
            stack.extend(adjacency[current] - component)
        unseen -= component
        component_slacks.append(slack)
    slack = component_slacks[0]
    model = NetworkModelInput(
        list(buses.values()),
        branches,
        transformers,
        [],
        [
            {"id": f"50hertz-boundary-{index}", "bus": bus, "kind": "external_grid", "vm_pu": 1.0}
            for index, bus in enumerate(component_slacks, start=1)
        ],
        [],
        [],
        slack,
        2025,
        {
            "source_url": SOURCE_PAGE,
            "license": LICENCE,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            "limitations": list(OMISSIONS),
        },
        model_id="50hertz-static-2025-q4",
        model_version="2025-Q4",
        validation_class="public_screening",
    )
    validate_network_model(model)
    workbook.close()
    quality = {
        "bus_count": len(buses),
        "line_count": len(branches),
        "transformer_count": len(transformers),
        "connected_component_count": len(component_slacks),
        "assumptions": assumptions,
        "limitations": list(OMISSIONS),
        "boundary_candidates": boundary_candidates,
        "remedial_actions": remedial_actions,
    }
    return model, quality


def write_normalized(model: NetworkModelInput, quality: dict[str, Any], data_root: Path) -> Path:
    digest = str(model.provenance["sha256"])
    destination = data_root / "normalized" / "networks" / "50hertz" / digest
    payload = {name: getattr(model, name) for name in model.__dataclass_fields__}
    atomic_write(
        destination / "network.json",
        (json.dumps(payload, indent=2, ensure_ascii=False) + "\n").encode(),
    )
    atomic_write(
        destination / "quality.json",
        (json.dumps(quality, indent=2, ensure_ascii=False) + "\n").encode(),
    )
    for file_name, key in (
        ("boundary-candidates.json", "boundary_candidates"),
        ("remedial-actions.json", "remedial_actions"),
    ):
        atomic_write(
            destination / file_name,
            (json.dumps(quality.get(key, []), indent=2, ensure_ascii=False) + "\n").encode(),
        )
    return destination
