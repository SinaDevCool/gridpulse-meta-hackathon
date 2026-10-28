"""Source-specific profiling for public Berlin and Bremen regional load series."""

from __future__ import annotations

from collections import Counter
from datetime import datetime, time
from io import BytesIO
from typing import Any

import openpyxl


def profile_berlin_load_workbook(
    content: bytes, *, source_id: str, snapshot_hash: str
) -> dict[str, Any]:
    workbook = openpyxl.load_workbook(BytesIO(content), data_only=True, read_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        declared_unit = "kW" if any(row and row[0] == "Max in kW" for row in rows[:10]) else None
        points = [
            row
            for row in rows
            if len(row) >= 3 and isinstance(row[0], datetime) and isinstance(row[1], time)
        ]
        values = [float(row[2]) for row in points if isinstance(row[2], (int, float))]
        sheets.append(
            {
                "sheet": sheet.title,
                "points": len(points),
                "declared_unit": declared_unit,
                "minimum_kw": min(values) if values else None,
                "maximum_kw": max(values) if values else None,
                "negative_values": sum(value < 0 for value in values),
                "missing_numeric_values": len(points) - len(values),
            }
        )
    return {
        "schema_version": "regional-load-quality-v1",
        "source_id": source_id,
        "snapshot_hash": snapshot_hash,
        "sheets": sheets,
        "total_points": sum(item["points"] for item in sheets),
        "units_validated": bool(sheets) and all(item["declared_unit"] == "kW" for item in sheets),
        "timezone_validated": False,
        "dst_semantics_validated": False,
        "geographic_resolution": "berlin_region_or_published_voltage_level",
        "feature_admitted": False,
        "warnings": [
            "Workbook dates and times do not carry UTC offsets.",
            "DST fold/gap semantics require operator documentation before UTC normalization.",
            "These are regional or voltage-level totals, not substation loading.",
        ],
    }


def profile_wesernetz_load_csv(
    content: bytes, *, source_id: str, snapshot_hash: str
) -> dict[str, Any]:
    text = content.decode("cp1252", errors="replace")
    points = []
    invalid = 0
    for line in text.splitlines():
        if ";" not in line:
            continue
        timestamp, value = line.split(";", 1)
        try:
            # Deliberately naive: the source does not publish an offset and stays quarantined.
            parsed = datetime.strptime(timestamp.strip(), "%d.%m.%Y %H:%M")  # noqa: DTZ007
            points.append((parsed, float(value.strip().replace(",", "."))))
        except ValueError:
            if timestamp[:1].isdigit():
                invalid += 1
    timestamps = [item[0] for item in points]
    counts = Counter(timestamps)
    values = [item[1] for item in points]
    return {
        "schema_version": "regional-load-quality-v1",
        "source_id": source_id,
        "snapshot_hash": snapshot_hash,
        "record_count": len(points),
        "coverage_local_start": min(timestamps).isoformat() if timestamps else None,
        "coverage_local_end": max(timestamps).isoformat() if timestamps else None,
        "duplicate_local_timestamps": sum(count - 1 for count in counts.values() if count > 1),
        "invalid_rows": invalid,
        "negative_values": sum(value < 0 for value in values),
        "minimum_reported_mw": min(values) if values else None,
        "maximum_reported_mw": max(values) if values else None,
        "units_validated": True,
        "timezone_validated": False,
        "dst_semantics_validated": False,
        "feature_admitted": False,
        "warnings": [
            "The file publishes local timestamps without explicit UTC offsets.",
            "Duplicate autumn DST times cannot be ordered until fold semantics are documented.",
            "The series is a regional aggregate and must not be allocated to substations.",
        ],
    }
