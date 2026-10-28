"""Load-series normalization with unit and timestamp safeguards."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from itertools import pairwise
from pathlib import Path
from typing import Any

import openpyxl

POWER_UNITS = {"MW": 1.0, "kW": 0.001}
APPARENT_POWER_UNITS = {"MVA": 1.0, "kVA": 0.001}
ENERGY_UNITS = {"MWh", "kWh", "GWh"}


@dataclass(frozen=True)
class LoadPoint:
    timestamp: datetime
    value: float
    unit: str
    voltage_level: str
    sheet: str
    row: int


@dataclass(frozen=True)
class LoadValidation:
    valid: bool
    errors: list[str]
    warnings: list[str]
    points: list[LoadPoint]


def normalize_power(value: Any, unit: str) -> float:
    if unit in ENERGY_UNITS:
        raise ValueError(f"energy unit {unit} cannot be normalized as instantaneous power")
    if unit in APPARENT_POWER_UNITS:
        raise ValueError(f"apparent-power unit {unit} cannot be silently converted to MW")
    if unit not in POWER_UNITS:
        raise ValueError(f"unsupported power unit: {unit}")
    return float(value) * POWER_UNITS[unit]


def parse_load_workbook(path: Path, *, voltage_level: str, unit: str = "MW") -> LoadValidation:
    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    errors: list[str] = []
    warnings: list[str] = []
    points: list[LoadPoint] = []
    seen: set[datetime] = set()
    for sheet in workbook.worksheets:
        for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if len(row) < 2 or not isinstance(row[0], datetime):
                continue
            try:
                value_mw = normalize_power(row[1], unit)
            except (TypeError, ValueError) as exc:
                errors.append(f"{sheet.title}!{row_number}: {exc}")
                continue
            timestamp = row[0]
            if timestamp in seen:
                errors.append(f"duplicate timestamp: {timestamp.isoformat()}")
                continue
            seen.add(timestamp)
            if value_mw < 0:
                errors.append(f"negative load: {sheet.title}!{row_number}")
                continue
            points.append(LoadPoint(timestamp, value_mw, "MW", voltage_level, sheet.title, row_number))
    points.sort(key=lambda item: item.timestamp)
    if not points:
        errors.append("no timestamp/value rows found")
    if len(points) > 1:
        gaps = [
            (right.timestamp - left.timestamp).total_seconds() / 60
            for left, right in pairwise(points)
        ]
        typical = min(gaps)
        if any(gap > typical * 1.5 for gap in gaps):
            warnings.append("time series contains missing intervals")
    return LoadValidation(not errors, errors, warnings, points)
